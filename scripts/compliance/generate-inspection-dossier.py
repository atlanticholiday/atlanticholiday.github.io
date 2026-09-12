from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "output" / "documents" / "Dossier_Inspecao_Registo_Tempos_Atlantic_Holiday.docx"

NAVY = "13213C"
TEAL = "0F766E"
SLATE = "475569"
LIGHT = "F1F5F9"
PALE_TEAL = "ECFDF5"
PALE_AMBER = "FFF7ED"
WHITE = "FFFFFF"
RED = "9F1239"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="CBD5E1", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, value, end])


def set_font(run, name="Arial", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_label_value(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    label_run = p.add_run(f"{label}: ")
    set_font(label_run, bold=True, color=NAVY)
    value_run = p.add_run(value)
    set_font(value_run, color=SLATE)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)
    return p


def add_note(doc, title, text, fill=PALE_TEAL, accent=TEAL):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, accent, "8")
    set_cell_margins(cell, 130, 150, 130, 150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(title)
    set_font(run, bold=True, color=accent)
    p = cell.add_paragraph(text)
    p.paragraph_format.space_after = Pt(0)
    return table


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            cell.width = widths[i]
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        set_font(run, size=8.5, bold=True, color=WHITE)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        cells = row.cells
        prevent_row_split(row)
        for i, value in enumerate(values):
            cell = cells[i]
            if row_index % 2:
                set_cell_shading(cell, "F8FAFC")
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if widths:
                cell.width = widths[i]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(value))
            set_font(run, size=8.2, color="1E293B")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = RGBColor.from_string("1E293B")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color, before, after in (
        ("Title", 29, NAVY, 0, 12),
        ("Subtitle", 12, SLATE, 0, 8),
        ("Heading 1", 17, NAVY, 14, 7),
        ("Heading 2", 12, TEAL, 10, 5),
        ("Heading 3", 10, NAVY, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(9.2)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("ATLANTIC HOLIDAY  |  REGISTO DE TEMPOS DE TRABALHO")
    set_font(run, size=7.5, bold=True, color=SLATE)

    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Cm(17.4))
    table.columns[0].width = Cm(13.8)
    table.columns[1].width = Cm(3.6)
    left = table.cell(0, 0).paragraphs[0]
    left.paragraph_format.space_after = Pt(0)
    run = left.add_run("Uso interno  •  Versão 1.0  •  12/09/2026")
    set_font(run, size=7.2, color=SLATE)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.paragraph_format.space_after = Pt(0)
    run = right.add_run("Página ")
    set_font(run, size=7.2, color=SLATE)
    add_page_number(right)


def build_document():
    doc = Document()
    configure_document(doc)

    # Cover
    doc.add_paragraph().paragraph_format.space_after = Pt(32)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("DOSSIER DE INSPEÇÃO")
    set_font(run, size=9, bold=True, color=TEAL)
    title = doc.add_paragraph(style="Title")
    title.add_run("Registo de tempos de trabalho")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Procedimento, evidências e plano de conformidade")

    add_note(
        doc,
        "Estado em 12 de setembro de 2026",
        "O sistema cobre os elementos técnicos centrais do registo de tempos, da rastreabilidade das correções, do trabalho suplementar e da consulta imediata. A conformidade final depende também da prática diária da empresa e das ações assinaladas como pendentes neste dossier.",
        PALE_AMBER,
        "B45309",
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(18)
    add_label_value(doc, "Entidade empregadora", "DAVID VASCONCELOS, UNIPESSOAL LDA")
    add_label_value(doc, "NIPC", "515 399 507")
    add_label_value(doc, "Estabelecimento", "Loja Atlantic Holiday, Caminho da Achada, Edifício Colinas da Achada, Bloco D, Porta 3.ª, R/C, Loja 1, 9000-208 Funchal")
    add_label_value(doc, "Atividade principal", "Atividades de serviços de intermediação de alojamento — CAE 55400-R4")
    add_label_value(doc, "Período de funcionamento", "Todos os dias, das 09:00 às 18:00")
    add_label_value(doc, "Contactos de privacidade", "+351 966 353 822 e +351 926 614 444")
    add_label_value(doc, "Responsável pela aprovação", "Gerência — aprovação interna pendente")
    doc.add_paragraph().paragraph_format.space_after = Pt(16)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Preparado para apoio à gestão e apresentação organizada de evidências perante a inspeção laboral")
    set_font(run, size=9, color=SLATE)

    doc.add_page_break()

    add_heading(doc, "1 Objetivo e conclusão", 1)
    doc.add_paragraph(
        "Este dossier descreve como a empresa regista os tempos de trabalho, conserva evidências e responde a pedidos de consulta. Foi preparado com base no Código do Trabalho, no RGPD, na Lei n.º 58/2019 e nos dados societários fornecidos pela empresa. Não substitui aconselhamento jurídico nem uma decisão da Autoridade para as Condições do Trabalho ou da Direção Regional do Trabalho da Madeira."
    )
    add_note(
        doc,
        "Conclusão operacional",
        "A aplicação está tecnicamente preparada para registar entrada, saída e intervalos, apurar tempos diários e semanais, conservar histórico auditável, imprimir ou exportar registos e tratar trabalho suplementar. Antes de considerar o processo encerrado, a empresa deve confirmar o IRCT/CCT aplicável, aprovar a informação de privacidade, afixar o mapa de horário, executar um teste real e instituir o arquivo mensal.",
    )

    add_heading(doc, "2 Identificação e organização do trabalho", 1)
    rows = [
        ("Firma", "DAVID VASCONCELOS, UNIPESSOAL LDA"),
        ("Natureza jurídica", "Sociedade unipessoal por quotas"),
        ("NIPC", "515 399 507"),
        ("Sede e local predominante", "Caminho da Achada, Edifício Colinas da Achada, Bloco D, Porta 3.ª, R/C, Loja 1, 9000-208 Funchal"),
        ("Atividade principal", "Atividades de serviços de intermediação de alojamento — CAE 55400-R4"),
        ("Período de funcionamento", "Todos os dias, 09:00–18:00"),
        ("Intervalo habitual", "Uma hora para refeição; o horário concreto consta do horário individual"),
        ("Descanso semanal", "Conforme o horário individual e a escala aplicável"),
        ("Dia de encerramento", "Sem dia fixo, de acordo com a informação fornecida"),
        ("IRCT/CCT", "Por confirmar por contabilista, consultor laboral ou Direção Regional do Trabalho"),
    ]
    add_table(doc, ["Elemento", "Informação de referência"], rows, [Cm(5.0), Cm(12.2)])
    doc.add_paragraph(
        "A certidão permanente consultada foi emitida em 2026 e indica, entre outras atividades, alojamento, intermediação, gestão, marketing, limpeza e serviços relacionados. O mapa de horário deve refletir a atividade efetivamente exercida no estabelecimento e os regimes individuais em vigor."
    )

    add_heading(doc, "3 Requisitos legais e evidência disponível", 1)
    legal_rows = [
        ("CT art. 202.º", "Registar início, termo e intervalos não incluídos; apurar horas por dia e semana; consulta imediata; conservar cinco anos.", "Picagens de entrada/saída/pausas, resumo diário e semanal, impressão, CSV e arquivo mensal verificável.", "Coberto tecnicamente"),
        ("CT arts. 215.º e 216.º", "Elaborar o mapa com identificação, atividade, local, período de funcionamento, horários, intervalos, descansos, IRCT e adaptabilidade; afixar em local visível.", "Configuração legal e mapa imprimível com horários individuais.", "Afixação e IRCT pendentes"),
        ("CT art. 217.º", "Consultar e avisar antes de alterações de horário, respeitando os prazos legais e acordos individuais.", "Histórico de horários e escalas na aplicação.", "Procedimento interno necessário"),
        ("CT art. 231.º", "Registar trabalho suplementar antes do início e após o termo, fundamento, visto do trabalhador e descansos compensatórios; consulta e impressão imediatas.", "Autorização, motivo, início, fim, validação do trabalhador, revisão da gestão e exportação anual.", "Coberto; exige uso disciplinado"),
        ("RGPD arts. 5.º, 13.º, 15.º, 16.º, 24.º, 25.º e 32.º", "Finalidade, transparência, minimização, acesso, retificação, responsabilidade e segurança adequadas ao risco.", "Acesso por função, conta pessoal, exportação dos próprios dados, pedido de correção, PIN não consultável e regras Firestore.", "Aviso formal pendente"),
        ("Lei n.º 58/2019, art. 28.º", "Tratamento de dados de trabalhadores para gestão laboral dentro dos limites legais e com deveres de sigilo.", "Separação entre dados pessoais, diretório mínimo, registos de assiduidade e acessos de gestão.", "Coberto tecnicamente"),
    ]
    add_table(doc, ["Base", "Obrigação", "Evidência do sistema", "Estado"], legal_rows, [Cm(2.5), Cm(5.1), Cm(6.2), Cm(3.4)])

    add_heading(doc, "4 Funcionamento do sistema", 1)
    add_heading(doc, "4.1 Identificação do trabalhador", 2)
    doc.add_paragraph(
        "Cada trabalhador pode marcar o próprio tempo através da conta pessoal. No tablet partilhado, a pessoa introduz apenas o seu PIN de seis algarismos; a estação identifica o perfil depois de validar o código. A lista completa de trabalhadores e os PINs não ficam expostos na estação."
    )
    add_bullet(doc, "O PIN é individual, não deve ser partilhado e deve ser alterado se houver suspeita de divulgação.")
    add_bullet(doc, "A gestão vê apenas se existe um PIN configurado; o valor não pode ser consultado na aplicação.")
    add_bullet(doc, "O tablet deve permanecer bloqueado à aplicação e sob controlo físico do estabelecimento.")

    add_heading(doc, "4.2 Eventos registados", 2)
    event_rows = [
        ("Entrada", "Início do período de trabalho", "Hora local de Portugal, hora UTC, identidade técnica e registo do servidor"),
        ("Início de pausa", "Início de intervalo não incluído no tempo de trabalho", "Evento autónomo e ordem cronológica"),
        ("Fim de pausa", "Retoma do trabalho", "Evento autónomo e ordem cronológica"),
        ("Saída", "Termo do período de trabalho", "Evento, total diário e estado concluído"),
        ("Correção manual", "Falha, esquecimento ou inexatidão", "Motivo obrigatório, autor, data, revisão e evento canónico"),
    ]
    add_table(doc, ["Evento", "Finalidade", "Rastreabilidade"], event_rows, [Cm(3.4), Cm(5.0), Cm(8.8)])
    doc.add_paragraph(
        "Os registos normais usam a hora do dispositivo e guardam também a hora UTC e a confirmação do servidor. As introduções manuais ficam identificadas como correções e exigem motivo e revisão. Nenhum utilizador pode apagar um registo de assiduidade; um evento incorreto é anulado por nova entrada auditada."
    )

    add_heading(doc, "4.3 Consulta e exportação", 2)
    add_bullet(doc, "Vista semanal e mensal por trabalhador, com total diário e total do período.")
    add_bullet(doc, "Impressão imediata do registo de tempos e exportação CSV.")
    add_bullet(doc, "Arquivo mensal JSON acompanhado de manifesto SHA-256 para verificar integridade.")
    add_bullet(doc, "Relação anual de trabalho suplementar em CSV.")
    add_bullet(doc, "Consulta pessoal e exportação dos próprios dados pelo trabalhador.")

    doc.add_page_break()
    add_heading(doc, "5 Procedimento diário", 1)
    for text in (
        "No início do trabalho, o trabalhador entra na própria conta ou introduz o PIN no tablet e confirma Entrada.",
        "Antes de um intervalo que não conta como tempo de trabalho, confirma Início de pausa; quando regressa, confirma Fim de pausa.",
        "No termo do trabalho, confirma Saída e verifica se o dia aparece como concluído.",
        "Se o trabalhador prestar atividade no exterior, envia ou valida o registo logo que regressa, assegurando que a empresa dispõe do visto dentro de 15 dias.",
        "Qualquer falha é comunicada no próprio dia. A gestão corrige por evento auditado, com motivo concreto, sem substituir nem apagar silenciosamente o original.",
    ):
        add_numbered(doc, text)

    add_heading(doc, "5.1 Falha de internet, energia ou tablet", 2)
    add_note(
        doc,
        "Regra de contingência",
        "Na indisponibilidade do sistema, usar o formulário do Anexo A. Assim que a aplicação voltar, a gestão introduz os eventos como correção manual, indica o motivo da indisponibilidade e conserva o formulário assinado juntamente com o arquivo do mês.",
        PALE_AMBER,
        "B45309",
    )
    add_bullet(doc, "Não inventar horas nem repetir picagens para compensar a falha.")
    add_bullet(doc, "Registar a hora observada, o local, o motivo e as assinaturas do trabalhador e de quem conferiu.")
    add_bullet(doc, "Se a sincronização ficar pendente, não limpar dados do navegador nem reinstalar a aplicação antes de confirmar o envio.")

    add_heading(doc, "6 Trabalho suplementar", 1)
    doc.add_paragraph(
        "O trabalho suplementar deve ser tratado num fluxo autónomo. Sempre que previsível, a gestão autoriza antes do início e regista o fundamento concreto, o local e a forma de compensação. O trabalhador marca o início e o fim, valida o registo e a gestão conclui a revisão."
    )
    overtime_rows = [
        ("Autorização", "Gestão", "Data, trabalhador, fundamento legal e factual, local, pagamento ou descanso"),
        ("Início", "Trabalhador", "Hora real imediatamente antes do trabalho suplementar"),
        ("Termo", "Trabalhador", "Hora real imediatamente após o fim"),
        ("Visto", "Trabalhador", "Validação após a prestação; para trabalho exterior, dentro do prazo aplicável"),
        ("Revisão", "Gestão", "Nota de conferência e data de descanso compensatório, quando aplicável"),
        ("Arquivo", "Gestão", "Relação anual e registos mensais conservados durante cinco anos"),
    ]
    add_table(doc, ["Momento", "Responsável", "Evidência mínima"], overtime_rows, [Cm(3.2), Cm(3.3), Cm(10.7)])
    add_note(
        doc,
        "Limites e compensação",
        "A aplicação documenta o processo, mas não substitui a conferência dos limites diários e anuais, do pagamento, dos descansos e das regras eventualmente mais favoráveis constantes do IRCT/CCT.",
        PALE_AMBER,
        "B45309",
    )

    add_heading(doc, "7 Correções, validação e não apagamento", 1)
    add_bullet(doc, "O trabalhador pode pedir correção a partir da própria conta e acompanhar o estado do pedido.")
    add_bullet(doc, "A gestão regista o motivo, o autor e a data de cada correção.")
    add_bullet(doc, "Os eventos originais mantêm-se; a anulação é representada por um novo evento canónico.")
    add_bullet(doc, "Os dias completos podem ser confirmados pelo trabalhador, mantendo histórico de confirmações.")
    add_bullet(doc, "Os registos com introdução manual ou inconsistência ficam numa fila de revisão.")

    doc.add_page_break()
    add_heading(doc, "8 Privacidade e acesso", 1)
    privacy_rows = [
        ("Finalidade", "Cumprir obrigações laborais, gerir assiduidade, apurar remuneração e descansos e responder a inspeções."),
        ("Base jurídica", "Cumprimento de obrigação jurídica e execução da relação laboral; não assentar no consentimento do trabalhador como regra geral."),
        ("Dados", "Identificação profissional, horário, picagens, pausas, correções, validações e trabalho suplementar."),
        ("Acesso", "Próprio trabalhador para os seus dados; gestão autorizada para funções laborais; estação apenas para autenticar e marcar."),
        ("Calendário de férias", "Visível à equipa apenas com informação operacional necessária. Não expor notas privadas, contactos, saldos, horas de ponto ou motivos de ausência."),
        ("Conservação", "Registos de tempos e trabalho suplementar durante cinco anos. Outros dados seguem o prazo necessário à respetiva finalidade ou obrigação."),
        ("Direitos", "Informação, acesso, retificação e demais direitos aplicáveis através dos contactos +351 966 353 822 e +351 926 614 444."),
        ("Prestadores", "Documentar os fornecedores que alojam ou tratam dados, as garantias contratuais e as localizações relevantes."),
    ]
    add_table(doc, ["Tema", "Regra interna"], privacy_rows, [Cm(4.0), Cm(13.2)])
    doc.add_paragraph(
        "A empresa deve entregar ou disponibilizar aos trabalhadores um aviso de privacidade aprovado, com a identidade do responsável, finalidades, bases jurídicas, categorias de destinatários, prazos, direitos e contacto. O facto de a aplicação mostrar esta informação não dispensa a aprovação interna do texto."
    )

    add_heading(doc, "9 Segurança e continuidade", 1)
    security_rows = [
        ("Contas pessoais", "Não partilhar palavras-passe; remover acesso quando termina a relação; rever funções trimestralmente."),
        ("Tablet", "Conta exclusiva de estação, ecrã bloqueado, acesso físico controlado, atualizações instaladas e sem aplicações pessoais."),
        ("PIN", "Seis algarismos únicos; entrega individual; alteração imediata em caso de suspeita; o valor não fica visível à gestão."),
        ("Firestore", "Regras impedem leitura cruzada dos dados pessoais e impedem alteração ou eliminação não autorizada de eventos."),
        ("Arquivo", "Exportar mensalmente dados e manifesto; guardar em dois locais separados com acesso restrito."),
        ("Incidente", "Registar perda, acesso indevido ou indisponibilidade; avaliar risco e obrigações de notificação RGPD."),
        ("Teste", "Testar trimestralmente uma restauração de arquivo e verificar o SHA-256 com o utilitário do projeto."),
    ]
    add_table(doc, ["Controlo", "Procedimento"], security_rows, [Cm(4.0), Cm(13.2)])

    doc.add_page_break()
    add_heading(doc, "10 Conteúdo do dossier de inspeção", 1)
    inspection_rows = [
        ("1", "Certidão permanente", "Certidão válida e identificação atualizada da entidade", "Disponível"),
        ("2", "Mapa de horário", "Versão atual, assinada ou aprovada e afixada; incluir horários individuais e escalas", "A concluir"),
        ("3", "Registos de tempos", "Impressão/CSV por trabalhador, dia, semana ou mês; últimos cinco anos", "Sistema pronto"),
        ("4", "Trabalho suplementar", "Autorizações, fundamentos, início/termo, vistos, descansos e relação anual", "Sistema pronto"),
        ("5", "Correções", "Pedido, motivo, evento original, correção, autor, data e revisão", "Sistema pronto"),
        ("6", "Arquivo de integridade", "JSON mensal e manifesto SHA-256 guardados separadamente", "Rotina a iniciar"),
        ("7", "Privacidade", "Aviso aos trabalhadores, matriz de acessos, lista de prestadores e registo de incidentes", "A aprovar"),
        ("8", "IRCT/CCT", "Identificação do instrumento aplicável ou parecer documentado de inexistência", "Por confirmar"),
        ("9", "Formação", "Lista de trabalhadores informados sobre marcação, correções e PIN", "A executar"),
        ("10", "Testes", "Relatório técnico e prova de um teste real controlado", "Técnico concluído; real pendente"),
    ]
    add_table(doc, ["N.º", "Documento", "Conteúdo", "Estado"], inspection_rows, [Cm(1.1), Cm(3.6), Cm(8.8), Cm(3.7)])

    add_heading(doc, "11 Plano de fecho", 1)
    plan_rows = [
        ("A", "Confirmar o IRCT/CCT e regras regionais aplicáveis", "Gerência com contabilista/consultor laboral ou DRT Madeira", "Antes da aprovação do mapa"),
        ("B", "Aprovar e afixar o mapa de horário; anexar escalas individuais", "Gerência", "Imediato"),
        ("C", "Aprovar e entregar o aviso de privacidade", "Gerência / responsável RGPD", "Antes do uso regular"),
        ("D", "Executar teste real com uma conta de gestão, uma conta de trabalhador e o tablet", "Gestão e trabalhador designado", "Antes da entrada em produção"),
        ("E", "Rever todas as picagens da primeira semana e corrigir desvios com motivo", "Gestão", "Primeira semana"),
        ("F", "Exportar o arquivo verificável e guardar cópia separada", "Gestão", "Mensal"),
        ("G", "Rever acessos, trabalhadores arquivados e PINs", "Administração", "Trimestral"),
        ("H", "Rever legislação, dossier e testes", "Gerência", "Anual ou após alteração relevante"),
    ]
    add_table(doc, ["Ação", "Medida", "Responsável", "Prazo"], plan_rows, [Cm(1.5), Cm(6.9), Cm(5.5), Cm(3.3)])

    add_heading(doc, "12 Migração do calendário de férias de 2026", 1)
    doc.add_paragraph(
        "A atualização preparada abrange 18 trabalhadores e usa o identificador técnico calendar-2026-v2. A fonte reconcilia 726 dias de direito acumulado, 185 dias usados até 2025, 270 dias reportados em 2026 e 271 dias remanescentes. O plano recusa a operação se faltar um trabalhador ou existir correspondência ambígua."
    )
    add_bullet(doc, "Correspondência prioritária pelo número de trabalhador e, em alternativa, pelo nome normalizado.")
    add_bullet(doc, "Cópia de segurança criada na mesma operação atómica, antes de substituir os registos de 2026.")
    add_bullet(doc, "Cópia limitada a saldos, baseline, férias e registos afetados; acesso reservado a funções de gestão.")
    add_bullet(doc, "Registos fora de 2026 são preservados, incluindo segmentos de períodos que atravessam a passagem de ano.")
    add_bullet(doc, "Aplicação única, identificada nas definições globais, e validação posterior de contagens e saldos.")
    add_note(
        doc,
        "Estado da execução",
        "A lógica, a cópia de segurança e os testes automáticos estão concluídos. A aplicação aos dados reais e a verificação visual exigem uma sessão autenticada de gestão; não são criadas picagens de assiduidade durante esta migração.",
        PALE_AMBER,
        "B45309",
    )

    add_heading(doc, "13 Evidência técnica", 1)
    tech_rows = [
        ("Testes da aplicação", "443 testes de navegador aprovados em 12/09/2026"),
        ("Regras de segurança", "Suite de integração Firestore aprovada, incluindo isolamento da cópia de férias"),
        ("Integridade", "Eventos canónicos imutáveis e arquivo mensal com manifesto SHA-256"),
        ("Retenção", "Campo retainUntil calculado por cinco anos para assiduidade e trabalho suplementar"),
        ("Versões anteriores", "fd775b6 — contacto RGPD no autosserviço; 26a2c17 — segurança do acesso pessoal a tarefas; f69b5ac — nomes completos na atualização de férias"),
    ]
    add_table(doc, ["Evidência", "Resultado"], tech_rows, [Cm(4.4), Cm(12.8)])

    doc.add_page_break()
    add_heading(doc, "Anexo A Registo manual de contingência", 1)
    doc.add_paragraph(
        "Usar apenas quando o registo digital estiver indisponível. Preencher no momento dos factos, sem rasuras não ressalvadas. Depois da reposição, transcrever por correção auditada e arquivar este original com o mês respetivo."
    )
    add_label_value(doc, "Data", "____ / ____ / ______")
    add_label_value(doc, "Trabalhador", "____________________________________________  N.º ______")
    add_label_value(doc, "Local", "____________________________________________________________")
    contingency_rows = [("", "", "", "", "") for _ in range(9)]
    add_table(
        doc,
        ["Entrada", "Início pausa", "Fim pausa", "Saída", "Observações"],
        contingency_rows,
        [Cm(2.7), Cm(3.0), Cm(3.0), Cm(2.7), Cm(5.8)],
    )
    add_label_value(doc, "Motivo da indisponibilidade", "____________________________________________________________")
    add_label_value(doc, "Assinatura do trabalhador", "____________________________________________________________")
    add_label_value(doc, "Conferido por", "________________________________  Data ____ / ____ / ______")
    add_label_value(doc, "Referência da correção digital", "____________________________________________________________")

    add_heading(doc, "Anexo B Lista de verificação mensal", 1)
    checklist_rows = [
        ("□", "Confirmar que todos os dias têm entrada, saída e intervalos aplicáveis."),
        ("□", "Rever dias incompletos, eventos manuais e anulações; documentar o resultado."),
        ("□", "Concluir validações dos trabalhadores e revisões de gestão pendentes."),
        ("□", "Conferir trabalho suplementar, fundamentos, compensações e descansos."),
        ("□", "Exportar registos e relação de trabalho suplementar quando aplicável."),
        ("□", "Criar arquivo JSON e manifesto SHA-256; guardar em locais separados."),
        ("□", "Confirmar que o mapa de horário e as escalas afixadas continuam atuais."),
        ("□", "Registar incidentes, alterações de acesso e saída de trabalhadores."),
    ]
    add_table(doc, ["Feito", "Controlo"], checklist_rows, [Cm(1.7), Cm(15.5)])
    add_label_value(doc, "Mês", "____________________")
    add_label_value(doc, "Responsável", "____________________________________________")
    add_label_value(doc, "Data e assinatura", "____________________________________________")

    doc.add_page_break()
    add_heading(doc, "Anexo C Fontes oficiais consultadas", 1)
    sources = [
        ("Código do Trabalho, Lei n.º 7/2009, versão consolidada", "https://diariodarepublica.pt/dr/legislacao-consolidada/lei/2009-34546475"),
        ("Código do Trabalho, artigo 202.º", "https://diariodarepublica.pt/dr/legislacao-consolidada/lei/2009-34546475-175393762"),
        ("Lei n.º 58/2019, execução do RGPD em Portugal", "https://diariodarepublica.pt/dr/detalhe/lei/58-2019-123815982"),
        ("Regulamento Geral sobre a Proteção de Dados", "https://eur-lex.europa.eu/eli/reg/2016/679/pt"),
        ("CNPD, biometria e assiduidade", "https://www.cnpd.pt/organizacoes/areas-tematicas/biometria/"),
        ("Jornal Oficial da Região Autónoma da Madeira, III Série", "https://joram.madeira.gov.pt/joram/3serie/"),
    ]
    add_table(doc, ["Fonte", "Endereço"], sources, [Cm(7.2), Cm(10.0)])
    doc.add_paragraph(
        "Consulta efetuada em 12 de setembro de 2026. As versões consolidadas facilitam a consulta, mas a validação jurídica deve considerar os atos publicados e o instrumento coletivo efetivamente aplicável."
    )

    add_heading(doc, "Aprovação interna", 1)
    add_label_value(doc, "Aprovado por", "____________________________________________________________")
    add_label_value(doc, "Função", "____________________________________________________________")
    add_label_value(doc, "Data", "____ / ____ / ______")
    add_label_value(doc, "Próxima revisão", "____ / ____ / ______")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
